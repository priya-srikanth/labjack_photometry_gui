"""Build standing PowerPoint decks from batch photometry or behavior figures."""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .config import load_analysis_config

NAVY = RGBColor(24, 52, 79)
GREY = RGBColor(90, 99, 108)


def _add_text(slide, text: str, top: float, size: int, bold: bool = False) -> None:
    frame = slide.shapes.add_textbox(Inches(.55), Inches(top), Inches(12.2), Inches(.8)).text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold = Pt(size), bold
    run.font.color.rgb = NAVY if bold else GREY


def build_figure_deck(root: Path, output: Path, title: str, pattern: str = "*.png") -> dict:
    """Create a reproducible figure deck; figure computation remains upstream."""
    figures = sorted(root.glob(f"sessions/*/*/*/{pattern}"))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    cover = deck.slides.add_slide(blank)
    _add_text(cover, title, 2.45, 38, True)
    _add_text(cover, f"{len(figures)} figures under {root}", 3.35, 16)
    for figure in figures:
        slide = deck.slides.add_slide(blank)
        relative = figure.relative_to(root)
        _add_text(slide, figure.stem.replace("_", " "), .14, 22, True)
        _add_text(slide, str(relative.parent), .72, 11)
        width_px, height_px = Image.open(figure).size
        max_w, max_h = Inches(12.5), Inches(6.05)
        width, height = max_w, int(max_w * height_px / width_px)
        if height > max_h:
            height, width = max_h, int(max_h * width_px / height_px)
        slide.shapes.add_picture(str(figure), int((deck.slide_width - width) / 2),
                                 Inches(1.2), width=width, height=height)
        slide.notes_slide.notes_text_frame.text = f"Source figure: {figure}"
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    return {"output": str(output), "slides": len(deck.slides), "figures": len(figures)}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("kind", choices=("photometry", "behavior"))
    parser.add_argument("--config", type=Path, default=Path("config/analysis.yaml"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    cfg = load_analysis_config(args.config)
    root = Path(cfg.output.photometry_root if args.kind == "photometry" else cfg.output.behavior_root)
    output = args.output or root / f"{args.kind}_summary_deck.pptx"
    pattern = "*.png" if args.kind == "photometry" else "lick_raster_by_position.png"
    result = build_figure_deck(root, output, f"{args.kind.capitalize()} summary", pattern)
    print(f"wrote {result['output']} ({result['slides']} slides, {result['figures']} figures)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
